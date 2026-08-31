import re
from pathlib import Path
from typing import List, Dict, Any

import pymupdf as fitz
import docx

from app.config import settings


class ProcessedChunk:
    def __init__(
        self,
        chunk_id: str,
        text: str,
        document_filename: str,
        category: str,
        filepath: str,
        page_number: int,
        chunk_index: int
    ):
        self.chunk_id = chunk_id
        self.text = text
        self.document_filename = document_filename
        self.category = category
        self.filepath = filepath
        self.page_number = page_number
        self.chunk_index = chunk_index

    def to_metadata(self) -> Dict[str, Any]:
        return {
            "chunk_id": self.chunk_id,
            "document_filename": self.document_filename,
            "category": self.category,
            "filepath": self.filepath,
            "page_number": int(self.page_number),
            "chunk_index": int(self.chunk_index)
        }


class DocumentProcessor:

    def __init__(
        self,
        chunk_size: int = None,
        chunk_overlap: int = None
    ):
        self.chunk_size = (
            chunk_size
            if chunk_size is not None
            else settings.CHUNK_SIZE
        )

        self.chunk_overlap = (
            chunk_overlap
            if chunk_overlap is not None
            else settings.CHUNK_OVERLAP
        )

        # Prevent invalid chunk settings
        if self.chunk_size <= 0:
            self.chunk_size = 1000

        if self.chunk_overlap < 0:
            self.chunk_overlap = 0

        if self.chunk_overlap >= self.chunk_size:
            self.chunk_overlap = self.chunk_size // 5

        self._rapid_ocr = None

    # ============================================================
    # TEXT CLEANING
    # ============================================================

    def clean_text(self, text: str) -> str:

        if not text:
            return ""

        text = text.replace("\xa0", " ")
        text = text.replace("\u200b", "")
        text = text.replace("\r\n", "\n")
        text = text.replace("\r", "\n")

        # Remove excessive spaces/tabs
        text = re.sub(r"[ \t]+", " ", text)

        # Remove excessive blank lines
        text = re.sub(r"\n\s*\n+", "\n\n", text)

        return text.strip()

    # ============================================================
    # OCR
    # ============================================================

    def _perform_ocr_on_page(self, page) -> str:
        """
        Try local RapidOCR for scanned PDF pages.

        OCR failure will not stop the entire ingestion process.
        """

        try:

            if self._rapid_ocr is None:
                from rapidocr_onnxruntime import RapidOCR
                self._rapid_ocr = RapidOCR()

            pix = page.get_pixmap(
                dpi=150,
                alpha=False
            )

            image_bytes = pix.tobytes("png")

            result, _ = self._rapid_ocr(image_bytes)

            if result:

                lines = []

                for item in result:

                    if (
                        isinstance(item, (list, tuple))
                        and len(item) >= 2
                        and item[1]
                    ):
                        lines.append(str(item[1]))

                text = "\n".join(lines).strip()

                if len(text) >= 15:
                    return text

        except Exception:
            # OCR is optional.
            pass

        return ""

    # ============================================================
    # PDF EXTRACTION
    # ============================================================

    def extract_text_from_pdf(
        self,
        pdf_path: str
    ) -> List[Dict[str, Any]]:

        pages_content = []

        try:

            with fitz.open(pdf_path) as pdf:

                for page_index in range(len(pdf)):

                    page = pdf[page_index]

                    text = page.get_text("text") or ""
                    cleaned = self.clean_text(text)

                    # OCR only if normal extraction failed
                    if len(cleaned) < 20:

                        ocr_text = self._perform_ocr_on_page(page)

                        if ocr_text:
                            cleaned = self.clean_text(ocr_text)

                    if cleaned:

                        pages_content.append(
                            {
                                "page_number": page_index + 1,
                                "text": cleaned
                            }
                        )

        except Exception as pdf_error:

            # Fallback to pypdf
            try:

                from pypdf import PdfReader

                reader = PdfReader(pdf_path)

                for page_index, page in enumerate(reader.pages):

                    text = page.extract_text() or ""
                    cleaned = self.clean_text(text)

                    if cleaned:

                        pages_content.append(
                            {
                                "page_number": page_index + 1,
                                "text": cleaned
                            }
                        )

            except Exception as fallback_error:

                raise RuntimeError(
                    f"Failed to extract PDF '{pdf_path}'. "
                    f"PyMuPDF error: {pdf_error}. "
                    f"Fallback error: {fallback_error}"
                )

        return pages_content

    # ============================================================
    # DOCX EXTRACTION
    # ============================================================

    def extract_text_from_docx(
        self,
        docx_path: str
    ) -> List[Dict[str, Any]]:

        pages_content = []

        try:

            document = docx.Document(docx_path)

            content = []

            # Paragraphs
            for paragraph in document.paragraphs:

                text = paragraph.text.strip()

                if text:
                    content.append(text)

            # Tables
            for table in document.tables:

                for row in table.rows:

                    cells = []

                    for cell in row.cells:

                        cell_text = cell.text.strip()

                        if cell_text:
                            cells.append(cell_text)

                    if cells:
                        content.append(" | ".join(cells))

            combined_text = "\n".join(content)

            cleaned = self.clean_text(combined_text)

            if cleaned:

                pages_content.append(
                    {
                        "page_number": 1,
                        "text": cleaned
                    }
                )

        except Exception as error:

            raise RuntimeError(
                f"Failed to extract DOCX '{docx_path}': {error}"
            )

        return pages_content

    # ============================================================
    # CHUNKING
    # ============================================================

    def split_text_into_chunks(
        self,
        text: str
    ) -> List[str]:

        if not text:
            return []

        text = text.strip()

        if len(text) <= self.chunk_size:
            return [text]

        chunks = []

        start = 0
        text_length = len(text)

        while start < text_length:

            end = min(
                start + self.chunk_size,
                text_length
            )

            # Last chunk
            if end >= text_length:

                chunk = text[start:end].strip()

                if chunk:
                    chunks.append(chunk)

                break

            split_position = -1

            # Prefer natural boundaries
            delimiters = [
                "\n\n",
                "\n",
                ". ",
                "? ",
                "! ",
                "; ",
                ", ",
                " "
            ]

            search_start = start + int(
                self.chunk_size * 0.60
            )

            for delimiter in delimiters:

                position = text.rfind(
                    delimiter,
                    search_start,
                    end
                )

                if position != -1:

                    split_position = (
                        position + len(delimiter)
                    )

                    break

            if (
                split_position <= start
                or split_position == -1
            ):
                split_position = end

            chunk = text[
                start:split_position
            ].strip()

            if chunk:
                chunks.append(chunk)

            # Overlap
            next_start = (
                split_position -
                self.chunk_overlap
            )

            if next_start <= start:
                next_start = split_position

            start = next_start

        return chunks

    # ============================================================
    # PROCESS FILE
    # ============================================================

    def process_file(
        self,
        file_path: str,
        category: str
    ) -> List[ProcessedChunk]:

        path = Path(file_path)

        if not path.exists():
            raise FileNotFoundError(
                f"File not found: {file_path}"
            )

        filename = path.name
        extension = path.suffix.lower()

        if extension == ".pdf":

            pages_data = self.extract_text_from_pdf(
                str(path)
            )

        elif extension == ".docx":

            pages_data = self.extract_text_from_docx(
                str(path)
            )

        else:

            return []

        safe_filename = re.sub(
            r"[^a-zA-Z0-9_-]",
            "_",
            filename
        )

        safe_category = re.sub(
            r"[^a-zA-Z0-9_-]",
            "_",
            category
        )

        all_chunks = []

        global_chunk_index = 0

        for page in pages_data:

            page_number = int(
                page["page_number"]
            )

            page_text = page["text"]

            page_chunks = (
                self.split_text_into_chunks(
                    page_text
                )
            )

            for local_index, chunk_text in enumerate(
                page_chunks
            ):

                if len(chunk_text.strip()) < 20:
                    continue

                chunk_id = (
                    f"{safe_category}_"
                    f"{safe_filename}_"
                    f"p{page_number}_"
                    f"c{local_index}"
                )

                chunk = ProcessedChunk(
                    chunk_id=chunk_id,
                    text=chunk_text,
                    document_filename=filename,
                    category=category,
                    filepath=str(path.resolve()),
                    page_number=page_number,
                    chunk_index=global_chunk_index
                )

                all_chunks.append(chunk)

                global_chunk_index += 1

        return all_chunks


document_processor = DocumentProcessor()